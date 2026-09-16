"""Import / sync the real student roster from the LJ spreadsheet.

Idempotent and NON-destructive — safe to re-run any time:
- adds a STUDENT account for each new enrollment number in the sheet,
- refreshes the name of students already present,
- leaves existing students (and all their submissions / scenarios / results) untouched.

Login identifier = enrollment number.  Password = last 7 digits of the enrollment number.

    python -m scripts.import_students                 # uses ../Lj student data.xlsx
    python -m scripts.import_students path/to/file.xlsx
    python -m scripts.import_students --prune         # also deactivate students no longer in the sheet
"""

from __future__ import annotations

import sys
from pathlib import Path

from sqlalchemy import select

from app.core.database import session_scope
from app.core.security import hash_password
from app.models.enums import Role
from app.models.user import User

DEFAULT_XLSX = Path(__file__).resolve().parents[2] / "Lj student data.xlsx"
SHEET = "Marks"


def read_roster(path: Path) -> list[tuple[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    out: list[tuple[str, str]] = []
    for row in list(ws.iter_rows(values_only=True))[1:]:  # skip header
        if not row or row[0] is None or row[1] is None:
            continue
        enrollment = str(row[0]).strip()
        name = str(row[1]).strip().title()
        if not enrollment.isdigit():
            print(f"  ! skipping non-numeric enrollment: {row!r}")
            continue
        out.append((enrollment, name))
    return out


def main() -> int:
    args = [a for a in sys.argv[1:] if a != "--prune"]
    prune = "--prune" in sys.argv
    path = Path(args[0]) if args else DEFAULT_XLSX
    if not path.exists():
        print(f"Spreadsheet not found: {path}", file=sys.stderr)
        return 1

    roster = read_roster(path)
    if not roster:
        print("No student rows found.", file=sys.stderr)
        return 1

    added = updated = 0
    with session_scope() as db:
        existing = {
            u.student_id: u
            for u in db.scalars(select(User).where(User.role == Role.STUDENT)).all()
        }
        seen: set[str] = set()
        for enrollment, name in roster:
            seen.add(enrollment)
            u = existing.get(enrollment)
            if u is None:
                db.add(
                    User(
                        name=name,
                        email=f"{enrollment}@lj.students.local",
                        password_hash=hash_password(enrollment[-7:]),
                        role=Role.STUDENT,
                        student_id=enrollment,
                        is_active=True,
                    )
                )
                added += 1
            else:
                if u.name != name or not u.is_active:
                    u.name, u.is_active = name, True
                    updated += 1

        deactivated = 0
        if prune:
            for sid, u in existing.items():
                if sid not in seen and u.is_active:
                    u.is_active = False
                    deactivated += 1

    print(f"Roster sync from {path.name}: {added} added, {updated} updated, "
          f"{len(existing) + added} total students"
          + (f", {deactivated} deactivated" if prune else ""))
    print(f"  Example: enrollment {roster[0][0]}  ·  password {roster[0][0][-7:]}  ({roster[0][1]})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
