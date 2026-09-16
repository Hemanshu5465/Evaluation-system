from __future__ import annotations

import io
import re

ALLOWED_SYLLABUS_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
}
ALLOWED_SYLLABUS_EXT = {".pdf", ".docx", ".txt"}


class UnsupportedDocument(Exception):
    pass


def extract_text(data: bytes, *, filename: str, content_type: str) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == ".pdf" or content_type == "application/pdf":
        return _pdf(data)
    if ext == ".docx" or "wordprocessingml" in content_type:
        return _docx(data)
    if ext == ".txt" or content_type.startswith("text/"):
        return data.decode("utf-8", errors="replace")
    raise UnsupportedDocument(f"Unsupported syllabus document: {filename} ({content_type})")


def _pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise UnsupportedDocument("PyMuPDF (fitz) is not installed; cannot extract PDF text.") from e
    text_parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def _docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise UnsupportedDocument("python-docx is not installed; cannot extract DOCX text.") from e
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


# --------------------------------------------------------------------------- #
# Heuristic structure extraction (used as a fallback / seed before AI runs)
# --------------------------------------------------------------------------- #
_TOPIC_LINE = re.compile(r"^\s*(?:unit|module|chapter|topic|week)\s*[:\-\d.]*\s*(.+)$", re.I)
_BULLET = re.compile(r"^\s*(?:[-*•]|\d+[.)])\s+(.+)$")
_OUTCOME_HINT = re.compile(r"(students? will|able to|understand|apply|analyse|analyze|design|implement|evaluate)", re.I)


def heuristic_structure(text: str) -> tuple[list[str], list[str]]:
    topics: list[str] = []
    outcomes: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or len(line) > 160:
            continue
        m = _TOPIC_LINE.match(line) or _BULLET.match(line)
        if m:
            candidate = m.group(1).strip(" .:-")
            if _OUTCOME_HINT.search(candidate):
                outcomes.append(candidate)
            elif 3 <= len(candidate) <= 90:
                topics.append(candidate)
    # de-dup preserving order
    topics = list(dict.fromkeys(topics))[:40]
    outcomes = list(dict.fromkeys(outcomes))[:20]
    return topics, outcomes
